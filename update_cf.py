import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import xlrd, openpyxl, shutil
from datetime import date
from collections import defaultdict
from openpyxl.styles import Font

# 1. 売掛一覧から4月発行分を入金予定日別に集計
wb_u = xlrd.open_workbook(r'\\YOUR_NAS_IP\YOUR_FOLDER2\01_billing\sales\accounts_receivable.xls')
ws_u = wb_u.sheet_by_name('Sheet1')

def xl_date(val):
    try:
        if val and val > 40000:
            return xlrd.xldate_as_datetime(val, wb_u.datemode).date()
    except:
        pass
    return None

groups = defaultdict(float)
for r in range(1, ws_u.nrows):
    try:
        hakkou = xl_date(ws_u.cell_value(r, 3))
        nyukin = xl_date(ws_u.cell_value(r, 5))
        kingaku = ws_u.cell_value(r, 4)
        if hakkou and date(2026,4,1) <= hakkou <= date(2026,4,30) and nyukin:
            groups[nyukin] += kingaku or 0
    except:
        pass

# 片桐さん(Just.Assist S2604-00027)は行11723で収入・外注費が別途計上済みのため除外
KITAGIRI = 2215400
may31_total   = int(groups.get(date(2026,5,31), 0))
may31_urikake = may31_total - KITAGIRI  # 行11721へ（SES含む全部まとめ）
jun10 = int(groups.get(date(2026,6,10), 0))
jun15 = int(groups.get(date(2026,6,15), 0))
jun20 = int(groups.get(date(2026,6,20), 0))
jun25 = int(groups.get(date(2026,6,25), 0))

print('=== 4月発行分 入金予定集計 ===')
print(f'  5月末 売掛合計         : {may31_total:>12,}円')
print(f'  5月末 片桐さん(行11723): {KITAGIRI:>12,}円 (除外・既計上)')
print(f'  5月末 売掛見込み(統合)  : {may31_urikake:>12,}円')
print(f'  6月10日                : {jun10:>12,}円')
print(f'  6月15日                : {jun15:>12,}円')
print(f'  6月20日                : {jun20:>12,}円')
print(f'  6月25日                : {jun25:>12,}円')

# 2. ファイルパス（バックアップから再作成）
CF_BACKUP = r'\\YOUR_NAS_IP\YOUR_FOLDER\01.cashflow\cashflow_backup\cashflow_YYYYMMDD.xlsx'
CF_NEW    = r'\\YOUR_NAS_IP\YOUR_FOLDER\01.cashflow\cashflow_new_YYYYMMDD.xlsx'

# 推測値読み取り用（data_only=Trueで計算済み値を取得）
wb_data = openpyxl.load_workbook(CF_BACKUP, data_only=True)
ws_data = wb_data['キャッシュ表_予定【融資有】']

# 編集用（数式保持）
wb_edit = openpyxl.load_workbook(CF_BACKUP)
ws_edit = wb_edit['キャッシュ表_予定【融資有】']

red_font = Font(color='FF0000')

def get_old_income(row_num):
    v = ws_data.cell(row_num, 5).value
    if isinstance(v, (int, float)) and v:
        return int(v), f'{int(v):,}'
    return 0, '(空)'

def update_income(row_num, new_amt, label):
    """E列(収入)を実績値に更新し、旧推測値をH列に赤文字で記載"""
    old_num, old_str = get_old_income(row_num)
    ws_edit.cell(row_num, 5).value = new_amt
    cell_h = ws_edit.cell(row_num, 8)
    cell_h.value = f'（推測値 {old_str}）'
    cell_h.font = Font(color='FF0000')
    diff = new_amt - old_num
    diff_str = f'+{diff:,}' if diff >= 0 else f'{diff:,}'
    print(f'  行{row_num} {label}')
    print(f'    E列: {old_str} → {new_amt:,}円  ({diff_str})')
    print(f'    H列: （推測値 {old_str}）[赤文字]')

def zero_with_note(row_num, note, label):
    """行を0クリアし、旧推測値とnoteをH列に赤文字で記載"""
    old_num, old_str = get_old_income(row_num)
    ws_edit.cell(row_num, 5).value = None
    cell_h = ws_edit.cell(row_num, 8)
    cell_h.value = f'{note}（旧推測値 {old_str}）'
    cell_h.font = Font(color='FF0000')
    print(f'  行{row_num} {label}')
    print(f'    E列: {old_str} → 空（{note}）')
    print(f'    H列: {note}（旧推測値 {old_str}）[赤文字]')

print('\n=== CF更新内容 ===')

# --- 5月末（月末） ---
print('\n[5月末]')
# 行11721: 売掛見込み → SES含む全月末分(片桐さん除く)をまとめて記載
update_income(11721, may31_urikake, '売掛見込み（SES統合）')
# 行11722: SES売掛見込み → 0クリア（上記に統合）
zero_with_note(11722, '→売掛見込みに統合', 'SES売掛見込み')

# --- 6月 ---
print('\n[6月]')
update_income(11744, jun10, '6月10日 売掛見込み')
update_income(11751, jun15, '6月15日 売掛見込み')
update_income(11757, jun20, '6月20日 売掛見込み')
update_income(11765, jun25, '6月25日 売掛見込み')

# 3. 保存
wb_edit.save(CF_NEW)
print(f'\n保存完了: {CF_NEW}')
